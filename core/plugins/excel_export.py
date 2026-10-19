"""plugins/excel_export.py — Экспорт вкладки «Свод» в Excel."""
__version__ = '1.0'
__description__ = 'Экспорт свода в Excel (сохранение фильтров, разметка, статусы)'

import io
from datetime import datetime, date as _date
from flask import request, send_file, jsonify

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False


# ============================================================
# Палитра
# ============================================================
C_HEADER_BG   = 'FF2C3E50'
C_HEADER_FG   = 'FFFFFFFF'
C_EMP_BG      = 'FFECF0F1'
C_TASK_BG     = 'FFFFFFFF'
C_TOTAL_BG    = 'FFFDEBD0'
C_TOTAL_FG    = 'FF7E5109'
C_WEEKEND_BG  = 'FFEEEEEE'
C_HOLIDAY_BG  = 'FFFFE0E0'
C_TODAY_BG    = 'FFC8E6C9'
C_TODAY_LINE  = 'FF2E7D32'
C_EMPTY_BG    = 'FFFAFAFA'
C_TITLE_FG    = 'FF1A1A1A'
C_SUB_FG      = 'FF555555'
C_META_FG     = 'FF888888'
C_BORDER      = 'FFCCCCCC'
C_FILTER_BG   = 'FFEAF2F8'

STATUS_FILL = {
    'done':        'FFC8E6C9',
    'in_progress': 'FFDCEDC8',
    'planned':     'FFE0E0E0',
    'blocked':     'FFFFCDD2',
}
STATUS_FG = {
    'done':        'FF1B5E20',
    'in_progress': 'FF33691E',
    'planned':     'FF424242',
    'blocked':     'FFB71C1C',
}
ABSENCE_FILL = {
    'vacation':      'FFBBDEFB',
    'sick_leave':    'FFFFE0B2',
    'business_trip': 'FFE1BEE7',
    'remote':        'FFB2EBF2',
}
ABSENCE_ICON = {
    'vacation':      '🏖',
    'sick_leave':    '🤒',
    'business_trip': '✈️',
    'remote':        '🏠',
}
ABSENCE_LABEL = {
    'vacation':      'Отпуск',
    'sick_leave':    'Больничный',
    'business_trip': 'Командировка',
    'remote':        'Удаленка',
}
STATUS_LABEL = {
    'done':        'Готово',
    'in_progress': 'В работе',
    'planned':     'Запланировано',
    'blocked':     'Заблокировано',
}
WD = ['Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс']


# ============================================================
# Стили
# ============================================================
def _fill(color):
    return PatternFill(fill_type='solid', fgColor=color)


def _font(bold=False, size=10, color=C_TITLE_FG):
    return Font(bold=bold, size=size, color=color)


def _center(wrap=False):
    return Alignment(horizontal='center', vertical='center', wrap_text=wrap)


def _left(wrap=True):
    return Alignment(horizontal='left', vertical='center', wrap_text=wrap)


def _thin(color=C_BORDER):
    s = Side(style='thin', color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _thick(color):
    s = Side(style='medium', color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _h(v):
    try:
        v = float(v)
    except Exception:
        return v
    if v == int(v):
        return int(v)
    return round(v, 2)


def _ddmm(iso):
    try:
        y, m, d = iso.split('-')
        return d + '.' + m
    except Exception:
        return iso


# PATCH 1.0.7 — маскирование ФИО в Excel (по флагу settings.excel_mask_names)
def _mask_name(name, enabled):
    if not enabled:
        return name
    s = str(name or '')
    if not s:
        return s
    return ' '.join(
        (w if len(w) <= 1 else w[0] + '*' * (len(w) - 1))
        for w in s.split()
    )


# ============================================================
# Lookup из БД (только для расшифровки фильтров)
# ============================================================
def _lookup(db):
    try:
        data = db.load_all()
    except Exception:
        return {'departments': {}, 'teams': {}, 'employees': {}, 'tasks': {}}
    return {
        'departments': {d['id']: d['name'] for d in data.get('departments', [])},
        'teams':       {t['id']: t['name'] for t in data.get('teams', [])},
        'employees':   {e['id']: e['name'] for e in data.get('employees', [])},
        'tasks':       {t['id']: t['name'] for t in data.get('tasks', [])},
    }


def _format_filters(f, lookup):
    parts = []
    if f.get('department'):
        parts.append('🏢 ' + (lookup['departments'].get(f['department']) or f['department']))
    if f.get('team'):
        parts.append('👥 ' + (lookup['teams'].get(f['team']) or f['team']))
    if f.get('employee'):
        parts.append('👤 ' + (lookup['employees'].get(f['employee']) or f['employee']))
    if f.get('task'):
        parts.append('📋 ' + (lookup['tasks'].get(f['task']) or f['task']))
    if f.get('tag'):
        tags = [x for x in f['tag'].split(',') if x]
        mode = 'И' if (f.get('tag_mode') or 'any').lower() == 'all' else 'ИЛИ'
        parts.append('🏷 ' + (' ' + mode + ' ').join('#' + t for t in tags))
    return ' · '.join(parts) if parts else 'без фильтров (только период)'


# ============================================================
# Лист «Свод»
# ============================================================
def _style_empty(cc, dk, day_meta, abs_map, today):
    m = day_meta.get(dk, {})
    abs_type = abs_map.get(dk)
    if abs_type:
        cc.value = ABSENCE_ICON.get(abs_type, '·')
        cc.fill = _fill(ABSENCE_FILL.get(abs_type, C_EMPTY_BG))
        cc.font = _font(size=11)
        cc.alignment = _center()
    elif m.get('is_holiday'):
        cc.value = '🎉'
        cc.fill = _fill(C_HOLIDAY_BG)
        cc.alignment = _center()
    elif m.get('is_weekend'):
        cc.fill = _fill(C_WEEKEND_BG)
    else:
        cc.fill = _fill(C_EMPTY_BG)


def _write_task_row(ws, r, t, days, col_first, col_total, day_meta, abs_map, today):
    icon = t.get('task_icon', '')
    name = t.get('task_name', '')
    c = ws.cell(r, 2, (icon + ' ' + name).strip())
    c.font = _font(bold=True, size=10, color='FFB71C1C') if t.get('overdue') else _font(size=10)
    c.alignment = _left()
    c.border = _thin()

    cells = t.get('cells') or {}
    for i, dk in enumerate(days):
        cc = ws.cell(r, col_first + i)
        cell = cells.get(dk)
        m = day_meta.get(dk, {})
        if cell and (cell.get('hours') or 0) > 0:
            cc.value = _h(cell['hours'])
            st = cell.get('status', 'done')
            cc.fill = _fill(STATUS_FILL.get(st, C_EMPTY_BG))
            cc.font = _font(bold=True, size=10, color=STATUS_FG.get(st, C_TITLE_FG))
            cc.alignment = _center()
            cc.border = _thin()
        else:
            _style_empty(cc, dk, day_meta, abs_map, today)

    tc = ws.cell(r, col_total, _h(t.get('total_hours', 0)))
    tc.font = _font(bold=True, size=10, color=C_TOTAL_FG)
    tc.fill = _fill(C_TOTAL_BG)
    tc.alignment = _center()
    tc.border = _thin()


def _build_main_sheet(ws, summary, meta, lookup):
    _mask = bool(meta.get('mask_names'))
    days = summary.get('days') or []
    rows = summary.get('rows') or []
    day_meta = summary.get('day_meta') or {}
    day_totals = summary.get('day_totals') or {}
    grand_total = summary.get('grand_total', 0)
    today = meta.get('today', '')

    n_days = len(days)
    col_emp = 1
    col_task = 2
    col_first = 3
    col_last = col_first + n_days - 1 if n_days else col_task
    col_total = (col_last + 1) if n_days else (col_task + 1)

    ws.cell(1, 1, 'Свод за период: ' + meta['from'] + ' — ' + meta['to']).font = _font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, col_total))
    ws.cell(1, 1).alignment = _left()
    ws.row_dimensions[1].height = 24

    ws.cell(2, 1, 'Фильтры: ' + _format_filters(meta['filters'], lookup)).font = _font(size=10, color=C_SUB_FG)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, col_total))
    ws.cell(2, 1).alignment = _left()
    ws.cell(2, 1).fill = _fill(C_FILTER_BG)
    ws.row_dimensions[2].height = 20

    ws.cell(3, 1, 'Сформировано: ' + meta['generated_at'] + ' · ' + meta['user']).font = _font(size=9, color=C_META_FG)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(2, col_total))
    ws.cell(3, 1).alignment = _left()
    ws.row_dimensions[3].height = 16

    hr = 5
    ws.row_dimensions[hr].height = 32
    hc = ws.cell(hr, col_emp, 'Сотрудник')
    hc.font = _font(bold=True, size=11, color=C_HEADER_FG)
    hc.fill = _fill(C_HEADER_BG)
    hc.alignment = _center()
    hc.border = _thin()

    hc = ws.cell(hr, col_task, 'Задача')
    hc.font = _font(bold=True, size=11, color=C_HEADER_FG)
    hc.fill = _fill(C_HEADER_BG)
    hc.alignment = _left()
    hc.border = _thin()

    for i, dk in enumerate(days):
        m = day_meta.get(dk, {})
        wd = WD[m.get('weekday', 0)] if 0 <= m.get('weekday', 0) <= 6 else ''
        label = _ddmm(dk) + '\n' + wd
        c = ws.cell(hr, col_first + i, label)
        c.font = _font(bold=True, size=10, color=C_HEADER_FG)
        c.fill = _fill(C_HEADER_BG)
        c.alignment = _center(wrap=True)
        c.border = _thin()
        if m.get('is_holiday'):
            c.fill = _fill(C_HOLIDAY_BG)
            c.font = _font(bold=True, size=10, color='FFB71C1C')
        elif m.get('is_weekend'):
            c.fill = _fill(C_WEEKEND_BG)
            c.font = _font(bold=True, size=10, color='FF555555')
        if dk == today or m.get('is_today'):
            c.fill = _fill(C_TODAY_BG)
            c.font = _font(bold=True, size=11, color='FF1B5E20')
            c.border = _thick(C_TODAY_LINE)

    hc = ws.cell(hr, col_total, 'Итого')
    hc.font = _font(bold=True, size=11, color=C_TOTAL_FG)
    hc.fill = _fill(C_TOTAL_BG)
    hc.alignment = _center()
    hc.border = _thin()

    r = hr + 1
    if not rows:
        ws.cell(r, 1, 'Нет данных за выбранный период').font = _font(size=11, color=C_META_FG)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, col_total))
        return r

    for emp in rows:
        tasks = emp.get('tasks') or []
        abs_map = emp.get('absence_days') or {}
        emp_name = _mask_name(emp.get('employee_name', ''), _mask)

        if not tasks:
            c = ws.cell(r, col_emp, emp_name)
            c.font = _font(bold=True, size=10)
            c.fill = _fill(C_EMP_BG)
            c.alignment = _left()
            c.border = _thin()
            c2 = ws.cell(r, col_task, '— нет задач —')
            c2.font = _font(size=10, color=C_META_FG)
            c2.alignment = _left()
            c2.border = _thin()
            for i, dk in enumerate(days):
                cc = ws.cell(r, col_first + i)
                cc.border = _thin()
                _style_empty(cc, dk, day_meta, abs_map, today)
            tc = ws.cell(r, col_total, 0)
            tc.font = _font(bold=True, size=10, color=C_TOTAL_FG)
            tc.fill = _fill(C_TOTAL_BG)
            tc.alignment = _center()
            tc.border = _thin()
            r += 1
            continue

        emp_row_start = r
        emp_row_end = r + len(tasks) - 1
        for ti, t in enumerate(tasks):
            _write_task_row(ws, emp_row_start + ti, t, days, col_first, col_total, day_meta, abs_map, today)
        c = ws.cell(emp_row_start, col_emp)
        c.value = emp_name
        c.font = _font(bold=True, size=10)
        c.fill = _fill(C_EMP_BG)
        c.alignment = _left()
        c.border = _thin()
        if emp_row_end > emp_row_start:
            ws.merge_cells(start_row=emp_row_start, start_column=col_emp, end_row=emp_row_end, end_column=col_emp)
            for k in range(emp_row_start, emp_row_end + 1):
                ws.cell(k, col_emp).fill = _fill(C_EMP_BG)
                ws.cell(k, col_emp).border = _thin()
        r = emp_row_end + 1

    rc = r
    ws.row_dimensions[rc].height = 22
    c = ws.cell(rc, 1, 'Итого за день')
    c.font = _font(bold=True, size=11, color=C_HEADER_FG)
    c.fill = _fill(C_HEADER_BG)
    c.alignment = _left()
    c.border = _thin()
    ws.merge_cells(start_row=rc, start_column=1, end_row=rc, end_column=col_task)
    ws.cell(rc, col_task).fill = _fill(C_HEADER_BG)
    ws.cell(rc, col_task).border = _thin()

    for i, dk in enumerate(days):
        m = day_meta.get(dk, {})
        v = day_totals.get(dk, 0)
        cc = ws.cell(rc, col_first + i, _h(v) if v else None)
        cc.font = _font(bold=True, size=10, color=C_TOTAL_FG)
        cc.fill = _fill(C_TOTAL_BG)
        cc.alignment = _center()
        cc.border = _thin()
        if m.get('is_holiday'):
            cc.fill = _fill(C_HOLIDAY_BG)
        elif m.get('is_weekend'):
            cc.fill = _fill(C_WEEKEND_BG)

    cc = ws.cell(rc, col_total, _h(grand_total))
    cc.font = _font(bold=True, size=12, color=C_TOTAL_FG)
    cc.fill = _fill(C_TOTAL_BG)
    cc.alignment = _center()
    cc.border = _thin()
    return rc


# ============================================================
# Лист «По сотрудникам»
# ============================================================
def _build_by_employee_sheet(ws, summary):
    rows = summary.get('rows') or []
    _mask = False
    try:
        _mask = bool(summary.get('_mask_names'))
    except Exception:
        _mask = False
    hdrs = ['ФИО', 'Отдел', 'Задач', 'Часов', 'Просроч.', 'Выполнено', 'В работе', '🏖 Отпуск', '🤒 Больничн.', '✈️ Команд.']
    widths = [30, 20, 10, 10, 11, 12, 12, 12, 13, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 1
    ws.cell(r, 1, 'Свод по сотрудникам').font = _font(bold=True, size=14)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(hdrs))
    r += 2

    hrow = r
    for i, h in enumerate(hdrs, 1):
        c = ws.cell(hrow, i, h)
        c.font = _font(bold=True, size=10, color=C_HEADER_FG)
        c.fill = _fill(C_HEADER_BG)
        c.alignment = _center(wrap=True)
        c.border = _thin()
    ws.row_dimensions[hrow].height = 30
    r += 1

    for emp in rows:
        tasks = emp.get('tasks') or []
        abs_map = emp.get('absence_days') or {}
        vac = sum(1 for v in abs_map.values() if v == 'vacation')
        sick = sum(1 for v in abs_map.values() if v == 'sick_leave')
        trip = sum(1 for v in abs_map.values() if v == 'business_trip')
        done = sum(1 for t in tasks if t.get('task_status') == 'done' or t.get('emp_completion_kind'))
        active = sum(1 for t in tasks if t.get('task_status') == 'active')
        overdue = sum(1 for t in tasks if t.get('overdue'))
        vals = [_mask_name(emp.get('employee_name', ''), _mask), emp.get('employee_department', ''),
                len(tasks), _h(emp.get('total_hours', 0)), overdue, done, active, vac, sick, trip]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border = _thin()
            if i == 1:
                c.font = _font(bold=True, size=10)
                c.alignment = _left()
            elif i == 2:
                c.font = _font(size=10)
                c.alignment = _left()
            elif i == 4:
                c.font = _font(bold=True, size=10, color=C_TOTAL_FG)
                c.fill = _fill(C_TOTAL_BG)
                c.alignment = _center()
            elif i == 5 and v:
                c.font = _font(bold=True, size=10, color='FFB71C1C')
                c.alignment = _center()
            elif i == 6:
                c.font = _font(size=10, color='FF1B5E20')
                c.alignment = _center()
            elif i == 7:
                c.font = _font(size=10, color='FF1565C0')
                c.alignment = _center()
            else:
                c.font = _font(size=10)
                c.alignment = _center()
        r += 1

    if rows:
        total_hours = sum(e.get('total_hours', 0) for e in rows)
        total_tasks = sum(len(e.get('tasks') or []) for e in rows)
        total_overdue = sum(1 for e in rows for t in (e.get('tasks') or []) if t.get('overdue'))
        total_done = sum(1 for e in rows for t in (e.get('tasks') or []) if t.get('task_status') == 'done' or t.get('emp_completion_kind'))
        total_active = sum(1 for e in rows for t in (e.get('tasks') or []) if t.get('task_status') == 'active')
        tot = ['ИТОГО', '', total_tasks, _h(total_hours), total_overdue, total_done, total_active, '', '', '']
        for i, v in enumerate(tot, 1):
            c = ws.cell(r, i, v)
            c.font = _font(bold=True, size=11, color=C_TOTAL_FG)
            c.fill = _fill(C_TOTAL_BG)
            c.alignment = _left() if i == 1 else _center()
            c.border = _thin()


# ============================================================
# Лист «По задачам»
# ============================================================
def _build_by_task_sheet(ws, summary):
    rows = summary.get('rows') or []
    agg = {}
    for emp in rows:
        for t in (emp.get('tasks') or []):
            tid = t.get('task_id')
            if not tid:
                continue
            a = agg.setdefault(tid, {
                'name': t.get('task_name', ''),
                'icon': t.get('task_icon', ''),
                'tags': list(t.get('task_tags') or []),
                'status': t.get('task_status', 'active'),
                'deadline': t.get('deadline', ''),
                'hours': 0.0,
                'workers': 0,
                'overdue_any': False,
            })
            a['hours'] += float(t.get('total_hours') or 0)
            a['workers'] += 1
            if t.get('overdue'):
                a['overdue_any'] = True

    data = sorted(agg.values(), key=lambda x: -x['hours'])
    hdrs = ['Задача', 'Теги', 'Статус', 'Дедлайн', 'Часов', 'Исполнителей']
    widths = [50, 25, 16, 12, 10, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 1
    ws.cell(r, 1, 'Свод по задачам').font = _font(bold=True, size=14)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(hdrs))
    r += 2

    hrow = r
    for i, h in enumerate(hdrs, 1):
        c = ws.cell(hrow, i, h)
        c.font = _font(bold=True, size=10, color=C_HEADER_FG)
        c.fill = _fill(C_HEADER_BG)
        c.alignment = _center(wrap=True)
        c.border = _thin()
    ws.row_dimensions[hrow].height = 28
    r += 1

    for a in data:
        label = (a['icon'] + ' ' + a['name']).strip()
        tags_txt = ', '.join('#' + t for t in a['tags'])
        vals = [label, tags_txt, STATUS_LABEL.get(a['status'], a['status']),
                a['deadline'], _h(a['hours']), a['workers']]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border = _thin()
            if i == 1:
                c.font = _font(bold=True, size=10, color=('FFB71C1C' if a['overdue_any'] else C_TITLE_FG))
                c.alignment = _left()
            elif i == 2:
                c.font = _font(size=9, color='FF6A4C00')
                c.alignment = _left()
            elif i == 3:
                c.font = _font(size=10)
                c.alignment = _center()
                stat_fill = STATUS_FILL.get(a['status'])
                if stat_fill:
                    c.fill = _fill(stat_fill)
            elif i == 5:
                c.font = _font(bold=True, size=10, color=C_TOTAL_FG)
                c.fill = _fill(C_TOTAL_BG)
                c.alignment = _center()
            else:
                c.font = _font(size=10)
                c.alignment = _center()
        r += 1


# ============================================================
# Лист «Инфо»
# ============================================================
def _build_info_sheet(ws, summary, meta):
    _mask = bool(meta.get('mask_names'))
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 30

    r = 1
    ws.cell(r, 1, 'Дополнительная информация').font = _font(bold=True, size=14)
    r += 2
    ws.cell(r, 1, 'Период: ' + meta['from'] + ' — ' + meta['to']).font = _font(size=11)
    r += 1
    ws.cell(r, 1, 'Сформировано: ' + meta['generated_at']).font = _font(size=11)
    r += 1
    ws.cell(r, 1, 'Пользователь: ' + meta['user']).font = _font(size=11)
    r += 2

    top = summary.get('top_tasks') or []
    if top:
        ws.cell(r, 1, '🏆 Топ задач по часам').font = _font(bold=True, size=12)
        r += 1
        for i, h in enumerate(['Задача', 'Часов', 'Записей'], 1):
            c = ws.cell(r, i, h)
            c.font = _font(bold=True, size=10, color=C_HEADER_FG)
            c.fill = _fill(C_HEADER_BG)
            c.border = _thin()
            c.alignment = _center()
        r += 1
        for t in top:
            vals = [(t.get('task_icon', '') + ' ' + t.get('task_name', '')).strip(),
                    _h(t.get('hours', 0)), t.get('count', 0)]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.font = _font(size=10)
                c.border = _thin()
                c.alignment = _left() if i == 1 else _center()
                if i == 2:
                    c.font = _font(bold=True, size=10, color=C_TOTAL_FG)
                    c.fill = _fill(C_TOTAL_BG)
            r += 1
        r += 1

    vacs = summary.get('vacations') or []
    if vacs:
        ws.cell(r, 1, '🏖 Отпуска / Больничные / Командировки в периоде').font = _font(bold=True, size=12)
        r += 1
        for i, h in enumerate(['Сотрудник', 'Тип', 'Период', 'Комментарий'], 1):
            c = ws.cell(r, i, h)
            c.font = _font(bold=True, size=10, color=C_HEADER_FG)
            c.fill = _fill(C_HEADER_BG)
            c.border = _thin()
            c.alignment = _center()
        r += 1
        for v in vacs:
            vtype = v.get('type', 'vacation')
            period = (v.get('from', '') or '') + ' — ' + (v.get('to', '') or '')
            vals = [_mask_name(v.get('employee_name', ''), _mask), ABSENCE_LABEL.get(vtype, vtype),
                    period, v.get('comment', '')]
            for i, val in enumerate(vals, 1):
                c = ws.cell(r, i, val)
                c.font = _font(size=10)
                c.border = _thin()
                c.alignment = _left()
                if i == 2:
                    c.fill = _fill(ABSENCE_FILL.get(vtype, C_EMPTY_BG))
                    c.alignment = _center()
            r += 1


# ============================================================
# Сборка книги
# ============================================================
def _build_workbook(summary, meta, lookup):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Свод'

    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 44
    n_days = len(summary.get('days') or [])
    for i in range(n_days):
        ws.column_dimensions[get_column_letter(3 + i)].width = 6.5
    ws.column_dimensions[get_column_letter(3 + n_days)].width = 10

    _build_main_sheet(ws, summary, meta, lookup)
    ws.freeze_panes = 'C6'

    _build_by_employee_sheet(wb.create_sheet('По сотрудникам'), summary)
    _build_by_task_sheet(wb.create_sheet('По задачам'), summary)
    _build_info_sheet(wb.create_sheet('Инфо'), summary, meta)

    return wb


# ============================================================
# JS: кнопка на вкладке «Свод»
# ============================================================
_JS = r'''
(function(){
'use strict';
function buildQuery(){
    var parts = [];
    function add(k, v){ if (v) parts.push(encodeURIComponent(k) + '=' + encodeURIComponent(v)); }
    var from = document.getElementById('sumDateFrom');
    var to   = document.getElementById('sumDateTo');
    var emp  = document.getElementById('sumEmployeeFilter');
    var task = document.getElementById('sumTaskFilter');
    var dept = document.getElementById('sumDeptFilter');
    var team = document.getElementById('sumTeamFilter');
    add('from', from && from.value);
    add('to', to && to.value);
    add('employee', emp && emp.value);
    add('task', task && task.value);
    add('department', dept && dept.value);
    add('team', team && team.value);
    if (window.__trkSelectedTags && window.__trkSelectedTags.sumTagFilter) {
        var tags = window.__trkSelectedTags.sumTagFilter;
        if (tags.length) add('tag', tags.join(','));
    }
    if (window.__trkTagMode && window.__trkTagMode.sumTagFilter) {
        add('tag_mode', window.__trkTagMode.sumTagFilter);
    }
    return parts.join('&');
}
function injectBtn(){
    if (!window.REESTR_PERMS || !window.REESTR_PERMS['tab.summary']) return;
    var toolbar = document.querySelector('#panel-summary .sum-toolbar');
    if (!toolbar) return;
    if (document.getElementById('excelExportBtn')) return;
    var btn = document.createElement('button');
    btn.className = 'btn btn-secondary';
    btn.id = 'excelExportBtn';
    btn.title = 'Выгрузить свод в Excel (с учётом текущих фильтров)';
    btn.textContent = '📥 Excel';
    btn.addEventListener('click', function(){
        var q = buildQuery();
        var url = '/api/plugins/excel_export/summary.xlsx' + (q ? '?' + q : '');
        window.location.href = url;
    });
    var applyBtn = document.getElementById('sumApplyBtn');
    if (applyBtn && applyBtn.parentNode === toolbar) {
        toolbar.insertBefore(btn, applyBtn);
    } else {
        toolbar.appendChild(btn);
    }
}
function boot(){ setTimeout(injectBtn, 400); }
if (document.readyState === 'loading') {
    document.addEventListener('DOMContentLoaded', boot);
} else { boot(); }
document.addEventListener('click', function(e){
    if (e.target && e.target.closest && e.target.closest('.tracker-tab[data-tab="summary"]')) {
        setTimeout(injectBtn, 250);
    }
});
})();
'''


# ============================================================
# Регистрация
# ============================================================
def register(app, db, hooks):
    print('[excel_export] v' + __version__ + ' · ' + __description__)
    if not _HAS_OPENPYXL:
        print('[excel_export] x openpyxl не установлен — экспорт в Excel недоступен')
        print('[excel_export]   Установите: pip install openpyxl')
        print('[excel_export]   Кнопка появится, но при клике вернёт 503')
    # PATCH v1.0.5: JS-роут регистрируется ВСЕГДА, чтобы не было 404 в браузере

    @app.route('/plugins-assets/excel_export.js', methods=['GET'])
    def _excel_export_js():
        return app.response_class(_JS, mimetype='application/javascript')

    @app.route('/api/plugins/excel_export/summary.xlsx', methods=['GET'])
    def _excel_export_summary():
        # PATCH v1.0.5: явный ответ вместо 500, если openpyxl нет
        if not _HAS_OPENPYXL:
            return jsonify({
                'success': False,
                'error': 'openpyxl не установлен на сервере. Выполните: pip install openpyxl'
            }), 503
        view = app.view_functions.get('api_summary_grid')
        if not view:
            return jsonify({'success': False, 'error': 'api_summary_grid not found'}), 500

        try:
            result = view()
        except Exception as e:
            import traceback
            traceback.print_exc()
            return jsonify({'success': False, 'error': 'summary failed: ' + str(e)}), 500

        if isinstance(result, tuple):
            body = result[0]
            code = result[1] if len(result) > 1 else 200
            if code != 200:
                return jsonify({'success': False, 'error': 'summary http ' + str(code)}), code
            result = body

        data = result.get_json(silent=True) if hasattr(result, 'get_json') else None
        if not isinstance(data, dict) or not data.get('success'):
            return jsonify({'success': False, 'error': 'summary not available'}), 401

        today_iso = _date.today().strftime('%Y-%m-%d')
        filters = {
            'department': (request.args.get('department') or '').strip(),
            'team':       (request.args.get('team') or '').strip(),
            'employee':   (request.args.get('employee') or '').strip(),
            'task':       (request.args.get('task') or '').strip(),
            'tag':        (request.args.get('tag') or '').strip(),
            'tag_mode':   (request.args.get('tag_mode') or 'any').strip().lower(),
        }

        try:
            from auth import current_user as _cur
            u = _cur(db) or {}
            user_name = u.get('display_name') or u.get('username') or ''
        except Exception:
            user_name = ''

        try:
            _mask_flag = str(db.get_setting('excel_mask_names', '0')) == '1'
        except Exception:
            _mask_flag = False

        meta = {
            'from': data.get('from') or '',
            'to':   data.get('to') or '',
            'filters': filters,
            'today': today_iso,
            'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'user': user_name,
            'mask_names': _mask_flag,
        }
        lookup = _lookup(db)

        try:
            data['_mask_names'] = _mask_flag
            wb = _build_workbook(data, meta, lookup)
        except Exception as e:
            import traceback
            traceback.print_exc()
            return jsonify({'success': False, 'error': 'xlsx build failed: ' + str(e)}), 500

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        fname = 'summary_' + (meta['from'] or 'start') + '_' + (meta['to'] or 'end') + '.xlsx'
        return send_file(
            buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=fname,
            max_age=0,
        )

    def on_startup(**ctx):
        print('[excel_export] готов — кнопка «📥 Excel» появится на вкладке «Свод»')
    hooks.on('on_startup', on_startup)
